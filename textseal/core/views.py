import json
import io
import os
import uuid
import zipfile
from pathlib import Path

from django.conf import settings
from django.db import transaction
from django.http import JsonResponse
from django.shortcuts import get_object_or_404, render
from django.urls import reverse
from django.views.decorators.http import require_GET, require_POST
from docx import Document
from pypdf import PdfReader

from .ledger import LedgerError, append_block, record_is_in_valid_chain, validate_chain
from .models import LedgerBlock, TextRecord, VerificationRecord
from .watermark import WatermarkError, canonicalize, content_hash, embed, extract, valid_token


def index(request):
    return render(request, "core/index.html")


def _payload(request):
    if request.content_type and request.content_type.startswith("application/json"):
        try:
            return json.loads(request.body or "{}")
        except json.JSONDecodeError as exc:
            raise WatermarkError("Request body must contain valid JSON.") from exc
    return request.POST


def _text_from_request(request, data):
    uploaded = request.FILES.get("file")
    text = data.get("text", "")
    if uploaded:
        if uploaded.size > settings.MAX_UPLOAD_SIZE:
            raise WatermarkError("Uploads must be 1 MB or smaller.")
        suffix = Path(uploaded.name).suffix.lower()
        try:
            content = uploaded.read()
            if suffix in {".txt", ".md"}:
                text = content.decode("utf-8")
            elif suffix == ".pdf":
                if not content.startswith(b"%PDF-"):
                    raise WatermarkError("This PDF file is invalid or unreadable.")
                reader = PdfReader(io.BytesIO(content), strict=False)
                if reader.is_encrypted:
                    raise WatermarkError("Password-protected PDF files are not supported.")
                if len(reader.pages) > 100:
                    raise WatermarkError("PDF files are limited to 100 pages.")
                text = "\n".join(page.extract_text() or "" for page in reader.pages)
                if not text.strip():
                    raise WatermarkError("No readable text was found in this PDF. Scanned-image OCR is not supported.")
            elif suffix == ".docx":
                with zipfile.ZipFile(io.BytesIO(content)) as package:
                    entries = package.infolist()
                    if "word/document.xml" not in package.namelist():
                        raise WatermarkError("This DOCX file does not contain a readable document.")
                    if len(entries) > 1000 or sum(entry.file_size for entry in entries) > 10_000_000:
                        raise WatermarkError("This DOCX file contains too much expanded data to process safely.")
                document = Document(io.BytesIO(content))
                paragraphs = [paragraph.text for paragraph in document.paragraphs]
                paragraphs.extend(cell.text for table in document.tables for row in table.rows for cell in row.cells)
                text = "\n".join(paragraphs)
                if not text.strip():
                    raise WatermarkError("No readable text was found in this DOCX file.")
            else:
                raise WatermarkError("Supported uploads are .txt, .md, .pdf, and .docx files.")
        except UnicodeDecodeError as exc:
            raise WatermarkError("TXT and MD uploads must be UTF-8 encoded.") from exc
        except (zipfile.BadZipFile, OSError, ValueError) as exc:
            if isinstance(exc, WatermarkError):
                raise
            raise WatermarkError("The uploaded document is invalid or unreadable.") from exc
        except Exception as exc:
            if isinstance(exc, WatermarkError):
                raise
            raise WatermarkError("The uploaded document could not be read safely.") from exc
    if not isinstance(text, str):
        raise WatermarkError("Text must be a string.")
    if len(text) > 100_000:
        raise WatermarkError("Text is limited to 100,000 characters.")
    return text


def _provenance_data(record):
    try:
        block = record.ledger_block
        ledger_status = record_is_in_valid_chain(record)
        block_number = block.block_index
    except LedgerBlock.DoesNotExist:
        ledger_status = False
        block_number = None
    return {
        "text_id": str(record.text_id),
        "content_hash": record.content_hash,
        "model_info": record.model_info or "Not provided",
        "source_note": record.source_note or "Not provided",
        "canonical_length": record.canonical_length,
        "blockchain_reference": record.blockchain_reference,
        "ledger_status": ledger_status,
        "block_number": block_number,
        "created_at": record.created_at.isoformat(),
    }


@require_GET
def health(request):
    return JsonResponse({
        "status": "ok",
        "blockchain_mode": os.environ.get("BLOCKCHAIN_MODE", "local-tamper-evident-ledger"),
        "ledger_integrity": validate_chain(),
    })


@require_POST
def watermark_text(request):
    try:
        data = _payload(request)
        original = _text_from_request(request, data)
        watermarked_text, text_id, token = embed(original)
        cleaned = canonicalize(original)
        with transaction.atomic():
            record = TextRecord.objects.create(
                text_id=text_id,
                content_hash=content_hash(cleaned),
                watermark_token=token,
                model_info=str(data.get("model_info", "")).strip()[:160],
                source_note=str(data.get("source_note", "")).strip()[:240],
                canonical_length=len(cleaned),
                blockchain_reference="pending",
            )
            append_block(record)
        return JsonResponse({
            "status": "registered",
            "message": "Watermark embedded and provenance added to the tamper-evident local ledger.",
            "watermarked_text": watermarked_text,
            "provenance": _provenance_data(record),
            "receipt_url": request.build_absolute_uri(reverse("receipt", kwargs={"text_id": record.text_id})),
            "watermark_note": "The watermark is invisible. Keep this exact output for later verification.",
        }, status=201)
    except (WatermarkError, LedgerError) as exc:
        return JsonResponse({"error": str(exc)}, status=400)


def _audit(record, submitted_hash, found, result, detail):
    VerificationRecord.objects.create(
        verification_id=uuid.uuid4(), text_record=record, submitted_hash=submitted_hash,
        watermark_found=found, result=result, detail=detail,
    )


def _verification_result(submitted, audit=True):
    watermark = extract(submitted)
    submitted_hash = content_hash(submitted)

    def finish(payload, status, record, found, result, detail):
        if audit:
            _audit(record, submitted_hash, found, result, detail)
        return payload, status

    if not watermark:
        detail = "No TextSeal watermark was found in the submitted text."
        return finish({
            "result": "not_registered", "label": "Not registered", "watermark_found": False,
            "hash_matches": False, "detail": detail,
        }, 200, None, False, "not_registered", detail)
    if not valid_token(watermark.text_id, watermark.token):
        detail = "A hidden payload was found, but its signature is invalid."
        return finish({
            "result": "invalid_watermark", "label": "Watermark invalid", "watermark_found": True,
            "hash_matches": False, "detail": detail,
        }, 200, None, True, "invalid_watermark", detail)
    record = TextRecord.objects.filter(text_id=watermark.text_id).first()
    if not record:
        detail = "The watermark is valid but this server has no matching provenance record."
        return finish({
            "result": "not_registered", "label": "Not registered", "watermark_found": True,
            "hash_matches": False, "detail": detail,
        }, 200, None, True, "not_registered", "Watermark refers to an unavailable record.")
    if record.watermark_token != watermark.token:
        detail = "The extracted watermark does not match its registered record."
        return finish({
            "result": "invalid_watermark", "label": "Watermark invalid", "watermark_found": True,
            "hash_matches": False, "detail": detail, "provenance": _provenance_data(record),
        }, 200, record, True, "invalid_watermark", "Record token mismatch.")
    if not record_is_in_valid_chain(record):
        detail = "The provenance ledger is unavailable or inconsistent. TextSeal will not claim authenticity."
        return finish({
            "result": "verification_unavailable", "label": "Verification unavailable", "watermark_found": True,
            "hash_matches": False, "detail": detail, "provenance": _provenance_data(record),
        }, 503, record, True, "verification_unavailable", detail)
    hash_matches = record.content_hash == submitted_hash
    if hash_matches:
        result, label, detail = "authentic", "Authentic & unchanged", "Watermark, provenance record, and SHA-256 hash all match."
    else:
        result, label, detail = "modified", "Modified / tampered", "A registered watermark was found, but the canonical SHA-256 hash differs."
    return finish({
        "result": result, "label": label, "watermark_found": True,
        "hash_matches": hash_matches, "detail": detail, "provenance": _provenance_data(record),
    }, 200, record, True, result, detail)


@require_POST
def verify_text(request):
    try:
        data = _payload(request)
        submitted = _text_from_request(request, data)
        payload, status = _verification_result(submitted)
        return JsonResponse(payload, status=status)
    except WatermarkError as exc:
        _audit(None, "", False, "invalid_input", str(exc))
        return JsonResponse({"error": str(exc)}, status=400)


@require_POST
def evaluate_robustness(request):
    try:
        data = _payload(request)
        submitted = _text_from_request(request, data)
        watermark = extract(submitted)
        if not watermark or not valid_token(watermark.text_id, watermark.token):
            raise WatermarkError("Paste a valid registered TextSeal output to evaluate.")
        record = TextRecord.objects.filter(text_id=watermark.text_id, watermark_token=watermark.token).first()
        if not record:
            raise WatermarkError("This watermark has no matching registered record on this server.")

        visible = watermark.visible_text
        suffix = submitted[len(visible):]
        cases = [
            ("Exact copy", submitted, "authentic"),
            ("Whitespace and newline normalization", f"  {visible.replace(chr(10), chr(13) + chr(10))}  \r\n{suffix}", "authentic"),
            ("Punctuation edit", f"{visible}!{suffix}", "modified"),
            ("Zero-width watermark removal", "".join(char for char in submitted if char not in {"\u200b", "\u200c", "\u200d", "\u2060", "\ufeff"}), "not_registered"),
        ]
        results = []
        for name, candidate, expected in cases:
            outcome, _ = _verification_result(candidate, audit=False)
            results.append({
                "name": name,
                "expected": expected,
                "actual": outcome["result"],
                "passed": outcome["result"] == expected,
            })
        passed = sum(case["passed"] for case in results)
        return JsonResponse({"cases": results, "passed": passed, "total": len(results)})
    except WatermarkError as exc:
        return JsonResponse({"error": str(exc)}, status=400)


@require_GET
def receipt(request, text_id):
    record = get_object_or_404(TextRecord, text_id=text_id)
    try:
        block_number = record.ledger_block.block_index
        ledger_status = record_is_in_valid_chain(record)
    except LedgerBlock.DoesNotExist:
        block_number = None
        ledger_status = False
    return render(request, "core/receipt.html", {
        "text_id": record.text_id,
        "content_hash": record.content_hash,
        "created_at": record.created_at,
        "blockchain_reference": record.blockchain_reference,
        "block_number": block_number,
        "ledger_status": ledger_status,
    })


@require_GET
def provenance(request, text_id):
    return JsonResponse({"provenance": _provenance_data(get_object_or_404(TextRecord, text_id=text_id))})


@require_GET
def history(request):
    """Small, anonymized activity feed for the final review demonstration."""
    records = TextRecord.objects.select_related("ledger_block").all()[:20]
    verifications = VerificationRecord.objects.select_related("text_record").all()[:20]
    return JsonResponse({
        "ledger_integrity": validate_chain(),
        "registered_count": TextRecord.objects.count(),
        "verification_count": VerificationRecord.objects.count(),
        "records": [_provenance_data(record) for record in records],
        "verifications": [{
            "verification_id": str(item.verification_id),
            "text_id": str(item.text_record.text_id) if item.text_record else None,
            "result": item.result,
            "watermark_found": item.watermark_found,
            "created_at": item.created_at.isoformat(),
        } for item in verifications],
    })
